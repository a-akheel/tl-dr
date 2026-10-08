import sqlite3
import tempfile
import time
import unittest
from io import BytesIO
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app import app
from summary import getPercentage, summarize
from webScrape import getArticle, public_target

TEXT = (Path(__file__).resolve().parents[1] / "samples/customer-rollout.txt").read_text()


def pdf_bytes():
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                             NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    content = DecodedStreamObject()
    content.set_data(b'BT /F1 12 Tf 30 700 Td (The customer needs an integration. The pilot covers one team. The next step is a workshop.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(content)
    output = BytesIO()
    writer.write(output)
    output.seek(0)
    return output


class AppTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        app.config.update(TESTING=True, SECRET_KEY='test-key', DATABASE=str(Path(self.directory.name) / 'test.sqlite3'))
        self.client = app.test_client()
        self.client.get('/')

    def tearDown(self):
        self.directory.cleanup()

    def post(self, data, client=None, **kwargs):
        client = client or self.client
        with client.session_transaction() as session:
            data['csrf_token'] = session['csrf_token']
        return client.post('/', data=data, **kwargs)

    def test_paste_and_download(self):
        response = self.post({'text': TEXT, 'copy-size': '3'})
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Your brief is ready', response.data)
        self.assertIn(b'no-store', response.headers['Cache-Control'].encode())
        export = self.client.get('/download/')
        self.assertEqual(export.status_code, 200)
        self.assertIn('summary.docx', export.headers['Content-Disposition'])
        document = Document(BytesIO(export.data))
        self.assertEqual(document.paragraphs[0].text, 'Summary')
        self.assertEqual([p.text for p in document.paragraphs[1:]], summarize(TEXT, n=3))

    def test_document_content_is_not_in_cookie(self):
        response = self.post({'text': TEXT, 'copy-size': '3'})
        self.assertLess(len(response.headers.get('Set-Cookie', '')), 1000)
        with self.client.session_transaction() as session:
            self.assertEqual(set(session), {'csrf_token', 'summary_id'})

    def test_each_upload_format(self):
        doc = Document()
        doc.add_paragraph(TEXT)
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        for stream, filename in [(BytesIO(TEXT.encode()), 'brief.txt'), (output, 'brief.docx'), (pdf_bytes(), 'brief.pdf')]:
            with self.subTest(filename=filename):
                response = self.post({'file': (stream, filename), 'upload-size': '3'}, content_type='multipart/form-data')
                self.assertEqual(response.status_code, 200)
                self.assertIn(b'Your brief is ready', response.data)

    def test_url_does_not_depend_on_preview(self):
        with patch('app.getArticle', return_value=('Integration article', TEXT)):
            self.assertEqual(self.post({'url': 'https://example.com/article', 'url-size': '3'}).status_code, 200)
        export = Document(BytesIO(self.client.get('/download/').data))
        self.assertEqual(export.paragraphs[0].text, 'Integration article')
        self.post({'text': TEXT, 'copy-size': '3'})
        export = Document(BytesIO(self.client.get('/download/').data))
        self.assertEqual(export.paragraphs[0].text, 'Summary')

    def test_invalid_inputs(self):
        for data in [{'text': ' ', 'copy-size': '3'}, {'text': TEXT, 'copy-size': 'bad'},
                     {'text': TEXT, 'copy-size': '0'}, {'text': TEXT, 'copy-size': '31'},
                     {'text': 'x' * 100001, 'copy-size': '3'}, {}]:
            self.assertEqual(self.post(data).status_code, 302)
        response = self.post({'file': (BytesIO(b'bad'), 'brief.exe'), 'upload-size': '3'}, follow_redirects=True)
        self.assertIn(b'Unsupported file', response.data)

    def test_corrupt_document_is_friendly_error(self):
        response = self.post({'file': (BytesIO(b'bad'), 'brief.docx'), 'upload-size': '3'}, follow_redirects=True)
        self.assertIn(b'Could not read this file', response.data)

    def test_oversized_upload(self):
        original = app.config['MAX_CONTENT_LENGTH']
        try:
            app.config['MAX_CONTENT_LENGTH'] = 1000
            response = self.post({'text': 'x' * 2000, 'copy-size': '3'}, follow_redirects=True)
            self.assertIn(b'upload is too large', response.data)
        finally:
            app.config['MAX_CONTENT_LENGTH'] = original

    def test_session_isolation(self):
        self.post({'text': TEXT, 'copy-size': '3'})
        other = app.test_client()
        self.assertEqual(other.get('/download/').status_code, 302)
        self.assertEqual(self.client.get('/download/').status_code, 200)

    def test_expiry(self):
        self.post({'text': TEXT, 'copy-size': '3'})
        with closing(sqlite3.connect(app.config['DATABASE'])) as connection:
            connection.execute('UPDATE summaries SET expires = ?', (time.time() - 1,))
            connection.commit()
        self.assertEqual(self.client.get('/download/').status_code, 302)
        with closing(sqlite3.connect(app.config['DATABASE'])) as connection:
            self.assertEqual(connection.execute('SELECT COUNT(*) FROM summaries').fetchone()[0], 0)

    def test_csrf(self):
        self.assertEqual(self.client.post('/', data={'text': TEXT, 'copy-size': '3'}).status_code, 302)
        self.assertEqual(self.client.get('/download/').status_code, 302)

    def test_summary_escapes_html(self):
        with patch('app.summarize', return_value=['<script>alert(1)</script>']):
            response = self.post({'text': TEXT, 'copy-size': '3'})
        self.assertIn(b'&lt;script&gt;', response.data)
        self.assertNotIn(b'<script>alert(1)</script>', response.data)


class SummaryTests(unittest.TestCase):
    def test_original_order_and_wording(self):
        result = summarize(TEXT, n=3)
        self.assertEqual(len(result), 3)
        self.assertTrue(all(sentence in TEXT for sentence in result))
        offsets = [TEXT.index(sentence) for sentence in result]
        self.assertEqual(offsets, sorted(offsets))

    def test_empty_and_short_text(self):
        self.assertEqual(summarize(''), [])
        self.assertEqual(summarize('One short sentence.', n=7), ['One short sentence.'])
        self.assertEqual(getPercentage('', []), 0)
        self.assertEqual(getPercentage('One short sentence.', ['One short sentence.']), 0)
        self.assertEqual(len(summarize(TEXT, title='the and', n=3)), 3)


class ArticleTests(unittest.TestCase):
    def test_blocks_private_targets(self):
        for address in ['127.0.0.1', '10.0.0.1', '169.254.169.254', '::1']:
            with patch('webScrape.socket.getaddrinfo', return_value=[(0, 0, 0, '', (address, 80))]):
                with self.assertRaises(ValueError):
                    public_target('http://example.com/article')

    def test_blocks_unsafe_schemes_credentials_and_ports(self):
        for url in ['file:///etc/passwd', 'http://user:pass@example.com', 'http://example.com:8080', 'nonsense']:
            with self.assertRaises(ValueError):
                public_target(url)

    def test_html_extraction(self):
        html = b'<html><title>Brief</title><nav><p>Menu</p></nav><article><p>First sentence.</p><p>Second sentence.</p></article></html>'
        with patch('webScrape.fetch_html', return_value=html):
            self.assertEqual(getArticle('https://example.com'), ('Brief', 'First sentence.\nSecond sentence.'))

    def test_redirect_revalidated(self):
        from unittest.mock import MagicMock
        from webScrape import fetch_html
        response = MagicMock(status=302, headers={'Location': 'http://127.0.0.1/'})
        pool = MagicMock()
        pool.request.return_value = response
        public = [(0, 0, 0, '', ('93.184.216.34', 443))]
        private = [(0, 0, 0, '', ('127.0.0.1', 80))]
        with patch('webScrape.socket.getaddrinfo', side_effect=[public, private]), patch('webScrape.urllib3.HTTPSConnectionPool', return_value=pool):
            with self.assertRaises(ValueError):
                fetch_html('https://example.com/article')
            pool.request.assert_called_once()
            response.close.assert_called_once()

    def test_dns_is_pinned(self):
        from unittest.mock import MagicMock
        from webScrape import fetch_html
        response = MagicMock(status=200, headers={'Content-Type': 'text/html'})
        response.stream.return_value = [b'<article><p>Hi.</p></article>']
        pool = MagicMock()
        pool.request.return_value = response
        public = [(0, 0, 0, '', ('93.184.216.34', 443))]
        with patch('webScrape.socket.getaddrinfo', return_value=public), patch('webScrape.urllib3.HTTPSConnectionPool', return_value=pool) as factory:
            self.assertIn(b'Hi.', fetch_html('https://example.com/article'))
            self.assertEqual(factory.call_args.args[0], '93.184.216.34')
            self.assertEqual(factory.call_args.kwargs['server_hostname'], 'example.com')
            self.assertEqual(pool.request.call_args.kwargs['headers']['Host'], 'example.com')


if __name__ == '__main__':
    unittest.main()
