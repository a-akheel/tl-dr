"""Local document summarizer. Cookies contain identifiers, never document contents."""

import io
import json
import os
import secrets
import sqlite3
import tempfile
import time
from pathlib import Path
from flask import Flask, abort, jsonify, render_template, request, send_file, session
from docx import Document
from summary import summarize, getPercentage, getTopWords
from readfile import getText
from webScrape import getArticle


def create_app(config=None):
    app = Flask(__name__)
    app.config.update(
        SECRET_KEY=os.environ.get("FLASK_SECRET_KEY") or secrets.token_hex(32),
        MAX_CONTENT_LENGTH=6 * 1024 * 1024,
        STATE_DB=os.environ.get(
            "STATE_DB", str(Path(app.instance_path) / "state.sqlite3")
        ),
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Strict",
    )
    app.config.update(config or {})
    db_path = Path(app.config["STATE_DB"])
    db_path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(db_path) as db:
        db.execute(
            "CREATE TABLE IF NOT EXISTS state (id TEXT PRIMARY KEY, payload TEXT, updated REAL)"
        )
    db_path.chmod(0o600)

    def state(value=None):
        sid = session.setdefault("sid", secrets.token_urlsafe(32))
        with sqlite3.connect(db_path) as db:
            db.execute("DELETE FROM state WHERE updated < ?", (time.time() - 86400,))
            if value is not None:
                db.execute(
                    "INSERT OR REPLACE INTO state VALUES (?, ?, ?)",
                    (sid, json.dumps(value), time.time()),
                )
                return value
            row = db.execute("SELECT payload FROM state WHERE id=?", (sid,)).fetchone()
            return json.loads(row[0]) if row else {}

    @app.before_request
    def protect_forms():
        session.setdefault("csrf", secrets.token_urlsafe(32))
        if request.method == "POST":
            token = request.form.get("csrf_token", "")
            if not secrets.compare_digest(token, session["csrf"]):
                abort(403)

    @app.after_request
    def private_response(response):
        response.headers["Cache-Control"] = "no-store"
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["X-Frame-Options"] = "DENY"
        return response

    @app.get("/")
    def index():
        return render_template("index.html")

    @app.post("/")
    def get_text():
        title = ""
        try:
            if "text" in request.form:
                text, size = request.form["text"], request.form.get("copy-size", "7")
            elif "url" in request.form:
                title, text = getArticle(request.form["url"])
                size = request.form.get("url-size", "7")
            else:
                upload = request.files.get("file")
                if not upload or not upload.filename:
                    abort(400, "Select a document.")
                suffix = Path(upload.filename).suffix.lower()
                if suffix not in {".txt", ".pdf", ".docx"}:
                    abort(400, "Supported files: TXT, PDF, DOCX.")
                with tempfile.TemporaryDirectory() as directory:
                    path = Path(directory) / ("document" + suffix)
                    upload.save(path)
                    text = getText(path)
                size = request.form.get("upload-size", "7")
            size = int(size)
            if not 1 <= size <= 30 or not text.strip() or len(text) > 100_000:
                abort(
                    400, "Enter 1–100,000 characters and a summary length from 1 to 30."
                )
            result = summarize(text, title, n=size)
            state({"title": title, "summary": result})
            return render_template(
                "result.html",
                summaries=result,
                percentage=getPercentage(text, result),
                words=getTopWords(text),
            )
        except (ValueError, UnicodeError):
            abort(400, "Invalid input or document.")
        except LookupError:
            abort(503, "Install the documented NLTK resources before summarizing.")

    @app.get("/download/")
    def download():
        current = state()
        if "summary" not in current:
            abort(404, "Generate a summary first.")
        doc = Document()
        doc.add_heading(current.get("title") or "Summary", level=2)
        for sentence in current["summary"]:
            doc.add_paragraph(sentence)
        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        return send_file(
            stream,
            download_name="Summary.docx",
            as_attachment=True,
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    @app.get("/_show_url_content/")
    def show_url_content():
        try:
            title, text = getArticle(request.args.get("url", ""))
            return jsonify(title=title, result=text)
        except ValueError:
            abort(400, "URL is unavailable or not explicitly allowed.")

    return app


app = create_app()
if __name__ == "__main__":
    app.run(host="127.0.0.1", debug=False)
